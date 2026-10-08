"""特定健診の平均値ブックを配信用 cube にする。

地域・年齢は第11回の都道府県別ブックだけ。二次医療圏・詳細情報・分布・質問票は読まない。
時代は第1–11回の同じブックから、全国の中計だけを抜く。第1回は年度末年齢。
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "ndb"
ERA_DIR = RAW / "era"
OUT = ROOT / "public" / "data"

# 健診年度。第1回（2013）から第11回（2023）。第12回に特定健診は無い。
ERA_YEARS = list(range(2013, 2024))
# 4行目の年齢見出し。区切りは全角チルダ（U+FF5E）。
AGE_BANDS = [
    "40～44歳",
    "45～49歳",
    "50～54歳",
    "55～59歳",
    "60～64歳",
    "65～69歳",
    "70～74歳",
    "中計",
]
# 平均値として置かれるハイフン。項目名の「γ-GT」とは別。
HYPHEN_CELLS = {"‐", "-", "－", "−", "–", "—", "―"}

# ブック上の表記（全角を含む）→ コード。
FILE_TO_CODE = {
    "BMI": "bmi",
    "腹囲": "waist",
    "空腹時血糖": "fpg",
    "HbA1C": "hba1c",
    "収縮期血圧": "sbp",
    "拡張期血圧": "dbp",
    "中性脂肪": "tg",
    "ＨＤＬ": "hdl",
    "ＬＤＬ": "ldl",
    "ＧＯＴ": "ast",
    "ＧＰＴ": "alt",
    "γ－ＧＴＰ": "gtp",
    "ヘモグロビン": "hb",
    "随時血糖": "cbg",
    "血清クレアチニン": "cre",
    "eGFR": "egfr",
}

ITEMS = [
    {"code": "bmi", "label": "BMI", "unit": "", "digits": 2},
    {"code": "waist", "label": "腹囲", "unit": "cm", "digits": 1},
    {"code": "fpg", "label": "空腹時血糖", "unit": "mg/dL", "digits": 1},
    {"code": "hba1c", "label": "HbA1C", "unit": "%", "digits": 2},
    {"code": "sbp", "label": "収縮期血圧", "unit": "mmHg", "digits": 1},
    {"code": "dbp", "label": "拡張期血圧", "unit": "mmHg", "digits": 1},
    {"code": "tg", "label": "中性脂肪", "unit": "mg/dL", "digits": 1},
    {"code": "hdl", "label": "HDL", "unit": "mg/dL", "digits": 1},
    {"code": "ldl", "label": "LDL", "unit": "mg/dL", "digits": 1},
    {"code": "ast", "label": "GOT", "unit": "U/L", "digits": 1},
    {"code": "alt", "label": "GPT", "unit": "U/L", "digits": 1},
    {"code": "gtp", "label": "γ-GTP", "unit": "U/L", "digits": 1},
    {"code": "hb", "label": "ヘモグロビン", "unit": "g/dL", "digits": 2},
    {"code": "cbg", "label": "随時血糖", "unit": "mg/dL", "digits": 1},
    {"code": "cre", "label": "血清クレアチニン", "unit": "mg/dL", "digits": 3},
    {"code": "egfr", "label": "eGFR", "unit": "mL/min/1.73m²", "digits": 1},
]

SEXES = [
    {"code": "male", "label": "男"},
    {"code": "female", "label": "女"},
]

AGES = [
    {"code": "a40", "label": "40–44歳"},
    {"code": "a45", "label": "45–49歳"},
    {"code": "a50", "label": "50–54歳"},
    {"code": "a55", "label": "55–59歳"},
    {"code": "a60", "label": "60–64歳"},
    {"code": "a65", "label": "65–69歳"},
    {"code": "a70", "label": "70–74歳"},
    {"code": "total", "label": "中計"},
]

# 列 3–18。男 7 階級 + 中計、女 7 階級 + 中計。
COLUMNS = [(sex["code"], age["code"]) for sex in SEXES for age in AGES]

MEAN_DIGITS = 4
REL_DIGITS = 4


def workbook_path() -> Path:
    found = [
        p
        for p in RAW.glob("*.xlsx")
        if "各項目の平均値" in p.name and "都道府県別" in p.name and "二次" not in p.name
    ]
    if len(found) != 1:
        raise SystemExit(f"平均値ブックが {len(found)} 件: {found}")
    return found[0]


def as_mean(value: object) -> float | None:
    """数値以外と、平均としてあり得ない 0 は欠測。埋め戻さない。"""
    if isinstance(value, str) and value.strip() in HYPHEN_CELLS:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if value == 0:
        return None
    return float(value)


# 第1–5回だけ長い名前。数値はあるので、後年の短い名前と同じコードにする。
ERA_LABELS = {
    "ＨＤＬコレステロール": "hdl",
    "ＬＤＬコレステロール": "ldl",
    "GOT（AST）": "ast",
    "GPT（ALT）": "alt",
    "γ-GT（γ-GTP）": "gtp",
}


def era_item_code(label: str) -> str | None:
    """第1–5回は単位つき、HbA1C は（NGSP）つきで表記が揺れる。コードへ寄せる。"""
    base = label.split("[", 1)[0].strip().replace("（NGSP）", "")
    if base.lower() == "hba1c":
        return "hba1c"
    aliased = ERA_LABELS.get(base)
    if aliased is not None:
        return aliased
    return FILE_TO_CODE.get(base)


def headers_match(header_rows: list[tuple[object, ...]]) -> bool:
    if len(header_rows) < 4:
        return False
    sex_row = list(header_rows[2]) + [None] * 18
    age_row = list(header_rows[3]) + [None] * 18
    if sex_row[2] != "男" or sex_row[10] != "女":
        return False
    expected = AGE_BANDS + AGE_BANDS
    actual = age_row[2:18]
    return actual == expected


def load_era_year(path: Path, year: int) -> dict[str, tuple[float | None, float | None]] | None:
    """全国の中計。見出しか年齢階級が第11回と違う年は None（値を作らない）。"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        if ws.title != "各項目の平均値":
            print(f"era {year}: シート名が違う {ws.title!r}。この年は null")
            return None
        header_rows: list[tuple[object, ...]] = []
        national: dict[str, tuple[float | None, float | None]] = {}
        pref: str | None = None
        saw_national = False
        for i, row in enumerate(ws.iter_rows(max_col=18, values_only=True), start=1):
            vals = tuple(row)
            if i <= 5:
                header_rows.append(vals)
                continue
            padded = list(vals) + [None] * (18 - len(vals))
            if padded[0]:
                pref = str(padded[0]).strip()
            if pref != "全国" or padded[1] is None:
                continue
            saw_national = True
            code = era_item_code(str(padded[1]).strip())
            if code is None:
                continue
            national[code] = (as_mean(padded[9]), as_mean(padded[17]))
        if not headers_match(header_rows):
            print(f"era {year}: 列見出しか年齢階級が違う。この年は null")
            return None
        if not saw_national:
            print(f"era {year}: 全国行が無い。この年は null")
            return None
        return national
    finally:
        wb.close()


def load_rows() -> tuple[list[dict[str, str]], dict[tuple[str, str], dict[str, float | None]]]:
    path = workbook_path()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    if ws.title != "各項目の平均値":
        raise SystemExit(f"シート名が違う: {ws.title}")

    areas: list[dict[str, str]] = [{"code": "00", "label": "全国"}]
    cells: dict[tuple[str, str], dict[str, float | None]] = {}
    pref_name: str | None = None
    pref_i = 0
    seen_items: set[str] = set()

    for row in ws.iter_rows(min_row=6, max_col=18, values_only=True):
        if row[0]:
            pref_name = str(row[0]).strip()
            if pref_name == "全国":
                code = "00"
            elif pref_name == "都道府県判別不可":
                code = "unk"
                areas.append({"code": code, "label": pref_name})
            else:
                pref_i += 1
                code = f"{pref_i:02d}"
                areas.append({"code": code, "label": pref_name})
        if pref_name is None or row[1] is None:
            continue
        label = str(row[1]).strip()
        item = FILE_TO_CODE.get(label)
        if item is None:
            raise SystemExit(f"未知の検査項目: {label}")
        seen_items.add(item)
        if pref_name == "全国":
            area = "00"
        elif pref_name == "都道府県判別不可":
            area = "unk"
        else:
            area = f"{pref_i:02d}"
        means = [as_mean(v) for v in row[2:18]]
        if len(means) != len(COLUMNS):
            raise SystemExit(f"{pref_name} {label} の列数が {len(means)}")
        cells[(item, area)] = {
            f"{sex}:{age}": mean for (sex, age), mean in zip(COLUMNS, means, strict=True)
        }

    expected = {item["code"] for item in ITEMS}
    if seen_items != expected:
        raise SystemExit(f"項目が揃わない: {sorted(expected - seen_items)}")
    if pref_i != 47:
        raise SystemExit(f"都道府県が {pref_i} 件")
    labels = [a["label"] for a in areas]
    if labels[0] != "全国" or "都道府県判別不可" not in labels or len(areas) != 49:
        raise SystemExit(f"地域が想定と違う: {labels}")
    return areas, cells


def build() -> None:
    areas, cells = load_rows()
    item_codes = [i["code"] for i in ITEMS]
    sex_codes = [s["code"] for s in SEXES]
    age_codes = [a["code"] for a in AGES]
    area_codes = [a["code"] for a in areas]

    nat = {item: cells[(item, "00")] for item in item_codes}
    size = len(item_codes) * len(sex_codes) * len(age_codes) * len(area_codes)
    mean: list[float | None] = [None] * size
    relative: list[float | None] = [None] * size

    def index(ii: int, si: int, ai: int, ari: int) -> int:
        return (
            ii * len(sex_codes) * len(age_codes) * len(area_codes)
            + si * len(age_codes) * len(area_codes)
            + ai * len(area_codes)
            + ari
        )

    for ii, item in enumerate(item_codes):
        for si, sex in enumerate(sex_codes):
            for ai, age in enumerate(age_codes):
                key = f"{sex}:{age}"
                base = nat[item][key]
                for ari, area in enumerate(area_codes):
                    value = cells[(item, area)][key]
                    at = index(ii, si, ai, ari)
                    if value is None:
                        continue
                    mean[at] = round(value, MEAN_DIGITS)
                    if base is None or base == 0:
                        continue
                    relative[at] = round(value / base, REL_DIGITS)

    dims = [
        {"name": "item", "codes": item_codes},
        {"name": "sex", "codes": sex_codes},
        {"name": "age", "codes": age_codes},
        {"name": "area", "codes": area_codes},
    ]
    dictionaries = {
        "items": ITEMS,
        "sexes": SEXES,
        "ages": AGES,
        "areas": areas,
        "year": 2023,
        "survey": "特定健診",
    }
    OUT.mkdir(parents=True, exist_ok=True)
    geo = {**dictionaries, "dims": dims, "measures": {"mean": mean, "relative": relative}}
    age = {**dictionaries, "dims": dims, "measures": {"mean": mean}}
    (OUT / "geo.json").write_text(
        json.dumps(geo, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    (OUT / "age.json").write_text(
        json.dumps(age, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    filled = sum(v is not None for v in mean)
    print(f"wrote geo.json age.json  cells={size} filled={filled}")
    build_era(item_codes, sex_codes)


def build_era(item_codes: list[str], sex_codes: list[str]) -> None:
    """item × sex × year。全国の中計だけ。男女は足さない。"""
    year_codes = [str(year) for year in ERA_YEARS]
    got: dict[int, dict[str, tuple[float | None, float | None]] | None] = {}
    for year in ERA_YEARS:
        path = ERA_DIR / f"{year}.xlsx"
        if not path.is_file():
            raise SystemExit(f"時代のブックが無い: {path}")
        got[year] = load_era_year(path, year)

    size = len(item_codes) * len(sex_codes) * len(year_codes)
    mean: list[float | None] = [None] * size

    def index(ii: int, si: int, yi: int) -> int:
        return ii * len(sex_codes) * len(year_codes) + si * len(year_codes) + yi

    for yi, year in enumerate(ERA_YEARS):
        national = got[year]
        if national is None:
            continue
        for ii, item in enumerate(item_codes):
            pair = national.get(item)
            if pair is None:
                continue
            for si, value in enumerate(pair):
                if value is None:
                    continue
                mean[index(ii, si, yi)] = round(value, MEAN_DIGITS)

    era = {
        "items": ITEMS,
        "sexes": SEXES,
        "years": [{"code": code, "label": f"{code}年度"} for code in year_codes],
        "survey": "特定健診",
        "age": "total",
        "area": "00",
        "dims": [
            {"name": "item", "codes": item_codes},
            {"name": "sex", "codes": sex_codes},
            {"name": "year", "codes": year_codes},
        ],
        "measures": {"mean": mean},
    }
    (OUT / "era.json").write_text(
        json.dumps(era, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    filled = sum(v is not None for v in mean)
    skipped = [year for year, national in got.items() if national is None]
    print(f"wrote era.json  cells={size} filled={filled} skipped={skipped}")


if __name__ == "__main__":
    build()
